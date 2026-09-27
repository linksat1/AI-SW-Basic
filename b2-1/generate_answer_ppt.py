"""평가자 질문사항 답변 자료(PDF) 생성.

질문마다 한 장: 요약 답변 + 코드 근거(파일:행 인용) + 실행 근거(명령어와 실제 출력).
코드 인용은 budget_app/*.py를 직접 읽어 행 번호를 붙이므로 코드가 바뀌면 다시 실행하면 된다.
실행 출력은 빈 demo 폴더에서 슬라이드 순서대로 실행해 얻은 실제 결과다.

사용법: python3 generate_answer_ppt.py  →  평가/B4_평가자_답변.pdf
"""
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from pptx_to_pdf import convert

BASE = Path(__file__).resolve().parent
OUT = BASE / "평가" / "B4_평가자_답변.pdf"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(18, 52, 98)
BLUE = RGBColor(24, 105, 182)
LIGHT = RGBColor(239, 246, 255)
WHITE = RGBColor(255, 255, 255)
TEXT = RGBColor(34, 34, 34)
MUTED = RGBColor(92, 102, 115)
GREEN = RGBColor(22, 140, 90)
ORANGE = RGBColor(196, 104, 20)
CODE_BG = RGBColor(30, 38, 52)
TERM_BG = RGBColor(12, 20, 16)
CODE_FG = RGBColor(226, 232, 240)
TERM_FG = RGBColor(190, 240, 200)
Q_FILL = RGBColor(255, 248, 230)
BORDER = RGBColor(215, 226, 239)

LEFT_X, RIGHT_X, COL_W = 0.72, 6.77, 5.85
LINE_H = 0.159  # 9pt 고정폭 글꼴 한 줄 높이(in)
BOTTOM = 7.0


def quote(file: str, *ranges) -> tuple[str, str]:
    """budget_app/<file>에서 행 범위를 읽어 (라벨, 행번호 붙은 코드)를 돌려준다."""
    lines = (BASE / "budget_app" / file).read_text(encoding="utf-8").splitlines()
    out, labels = [], []
    for r in ranges:
        a, b = (r, r) if isinstance(r, int) else r
        if out:
            out.append("    ⋮")
        out.extend(f"{n:>3}  {lines[n - 1]}" for n in range(a, b + 1))
        labels.append(f"{a}" if a == b else f"{a}-{b}")
    return f"{file}:{', '.join(labels)}", "\n".join(out)


def add_footer(slide):
    box = slide.shapes.add_textbox(Inches(11.95), Inches(7.1), Inches(0.7), Inches(0.2))
    p = box.text_frame.paragraphs[0]
    p.text = f"{len(prs.slides):02d}"
    p.alignment = PP_ALIGN.RIGHT
    p.font.size = Pt(9)
    p.font.color.rgb = MUTED


def text(slide, x, y, w, h, value, size, color=TEXT, bold=False, wrap=True, font=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        if font:
            p.font.name = font
    return box


def rect(slide, x, y, w, h, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
    else:
        shape.line.fill.background()


def new_slide(title, questions=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT
    text(slide, 0.72, 0.3, 11.9, 0.55, title, 24, NAVY, bold=True, wrap=False)
    add_footer(slide)
    if not questions:
        return slide, 1.05
    qs = questions if isinstance(questions, list) else [questions]
    h = 0.14 + 0.27 * sum(-(-len(q) // 90) for q in qs)
    rect(slide, 0.72, 0.92, 11.9, h, Q_FILL, RGBColor(240, 210, 150))
    text(slide, 0.85, 0.94, 11.6, h, "\n".join(f"Q. {q}" for q in qs), 12.5, ORANGE, bold=True)
    return slide, 0.92 + h + 0.12


def answer(slide, y, lines):
    """질문 아래 한두 줄 요약 답변. 다음 y를 돌려준다."""
    h = 0.16 + 0.26 * len(lines)
    rect(slide, 0.72, y, 11.9, h, WHITE, BORDER)
    body = "\n".join(f"A. {l}" if i == 0 else f"    {l}" for i, l in enumerate(lines))
    text(slide, 0.85, y + 0.03, 11.6, h, body, 13, TEXT)
    return y + h + 0.12


WRAP_COLS = 66  # 9pt Menlo로 상자 폭에 들어가는 글자 수 (한글은 2칸)


def wrap_body(body: str) -> str:
    """상자 폭을 넘는 줄을 직접 나눠, 줄 수로 계산한 상자 높이가 실제와 맞게 한다."""
    out = []
    for line in body.split("\n"):
        cur, width = "", 0
        for ch in line:
            w = 2 if ord(ch) > 0x2E80 else 1
            if width + w > WRAP_COLS:
                out.append(cur)
                cur, width = "      ", 6
            cur += ch
            width += w
        out.append(cur)
    return "\n".join(out)


def visual_lines(body: str) -> int:
    """화면에서 자동 줄바꿈된 뒤의 줄 수. 본문에는 줄바꿈을 넣지 않아 복사해도 명령이 한 줄로 유지된다."""
    return wrap_body(body).count("\n") + 1


def block_height(block) -> float:
    label, body, _ = block
    return (0.24 if label else 0) + 0.14 + visual_lines(body) * LINE_H


def column(slide, x, y, header, blocks, fresh=True):
    """blocks를 세로로 쌓는다. 하단을 넘는 블록들과 마지막 y를 돌려준다.
    fresh=True(빈 열)면 첫 블록은 넘치더라도 배치한다."""
    if not blocks or (not fresh and y + 0.3 + block_height(blocks[0]) > BOTTOM):
        return blocks, y
    text(slide, x, y, COL_W, 0.3, header, 12, GREEN if header.startswith("실행") else BLUE, bold=True, wrap=False)
    y += 0.3
    for i, block in enumerate(blocks):
        if (i > 0 or not fresh) and y + block_height(block) > BOTTOM:
            return blocks[i:], y
        label, body, kind = block
        if label:
            text(slide, x, y, COL_W, 0.25, label, 10, MUTED, bold=True, wrap=False)
            y += 0.24
        h = 0.14 + visual_lines(body) * LINE_H
        rect(slide, x, y, COL_W, h, CODE_BG if kind == "code" else TERM_BG)
        text(slide, x + 0.05, y + 0.02, COL_W - 0.1, h, body, 9, CODE_FG if kind == "code" else TERM_FG,
             font="Menlo")
        y += h + 0.1
    if y - 0.1 > BOTTOM + 0.05:
        print(f"[경고] {len(prs.slides)}번 슬라이드 '{header}' 블록 하나가 하단을 넘습니다 (y={y:.2f})")
    return [], y


def code(file, *ranges):
    label, body = quote(file, *ranges)
    return (label, body, "code")


def term(body, label=None):
    return (label, body, "term")


def qa(title, questions, lines, code_blocks, run_blocks, note=None):
    """질문 한 장. 코드·실행 근거가 넘치면 같은 질문으로 '(계속)' 슬라이드를 이어 붙인다."""
    slide, y = new_slide(title, questions)
    y = answer(slide, y, lines)
    if note:
        text(slide, 0.72, 7.02, 11.0, 0.3, note, 10, ORANGE, bold=True, wrap=False)
    code_hdr, run_hdr = "코드 근거 (budget_app/파일:행)", "실행 근거 (b2-1 폴더에서 실행)"
    while True:
        code_blocks, ly = column(slide, LEFT_X, y, code_hdr, code_blocks)
        run_blocks, ry = column(slide, RIGHT_X, y, run_hdr, run_blocks)
        # 한쪽 열이 넘치면 다른 열의 남은 공간부터 쓴다
        if code_blocks and not run_blocks:
            code_blocks, _ = column(slide, RIGHT_X, ry, code_hdr, code_blocks, fresh=False)
        elif run_blocks and not code_blocks:
            run_blocks, _ = column(slide, LEFT_X, ly, run_hdr, run_blocks, fresh=False)
        if not code_blocks and not run_blocks:
            return
        slide, y = new_slide(f"{title} (계속)", questions)


# ---------------------------------------------------------------------------
# 01 표지
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
text(slide, 0.9, 2.3, 11.5, 0.8, "B2-1 가계부 콘솔 프로그램", 40, WHITE, bold=True, wrap=False)
text(slide, 0.9, 3.2, 11.5, 0.6, "평가자 질문사항에 대한 답변", 24, RGBColor(190, 214, 245), bold=True, wrap=False)
text(slide, 0.9, 4.2, 11.5, 0.9,
     "질문마다  요약 답변  ·  코드 근거(파일:행 인용)  ·  실행 근거(명령어와 실제 출력)\n"
     "항목 1 기능·영속성   항목 2 구조·책임   항목 3 파이썬 기법   항목 4 설계 판단",
     15, RGBColor(160, 185, 220))

# 02 시연 준비
slide, y = new_slide("시연 준비 | 실행 환경과 순서")
column(slide, LEFT_X, y, "실행 근거 (b2-1 폴더에서 실행)", [
    term("$ cd b2-1\n"
         "$ cat > imp.csv <<'EOF'\n"
         "date,type,category,amount,memo,tags\n"
         "2024-01-20,expense,transport,1250,지하철,\n"
         "2024-01-21,expense,없는카테고리,1000,깨짐,\n"
         "2024-01-22,income,salary,50000,보너스,\n"
         "EOF\n"
         "$ cat imp.csv        # 헤더 + 3행 확인",
         "① 가져오기 테스트용 CSV 준비 (2번째 데이터 행은 일부러 깨진 행)"),
    term("$ python3 -m budget_app <명령> [옵션] --data-dir demo",
         "② 모든 명령은 빈 demo 폴더를 대상으로 실행"),
])
rect(slide, RIGHT_X, y + 0.3, COL_W, 5.5, WHITE, BORDER)
text(slide, RIGHT_X + 0.2, y + 0.45, COL_W - 0.4, 5.2,
     "읽는 법\n"
     "• 이후 슬라이드의 명령을 순서대로 실행하면 같은 출력이 나옵니다.\n"
     "• 거래 id(TX-000001…)는 빈 폴더 기준으로 매겨집니다.\n"
     "• 코드 근거의 왼쪽 숫자는 실제 파일의 행 번호입니다.\n\n"
     "주의: --data-dir는 명령 이름 뒤에\n"
     "• 맞음: python3 -m budget_app add --data-dir demo\n"
     "• 틀림: python3 -m budget_app --data-dir demo add\n"
     "  → 하위 명령의 기본값(data)이 앞의 값을 덮어써서\n"
     "     실제 data/ 폴더에 기록됩니다.\n\n"
     "명령 복사는 PDF 대신 평가/B4a_시연_명령어.sh에서\n"
     "• PDF는 긴 줄이 화면 폭에서 끊긴 채로 복사됩니다.\n"
     "• 전체 실행: zsh 평가/B4a_시연_명령어.sh\n\n"
     "시연 후 정리: rm -r demo imp.csv out.csv", 13, TEXT)

# ---------------------------------------------------------------------------
# 항목 1
qa("항목 1-1 | 8개 명령 동작 (1/2): add · list · search",
   "add/list/search/summary/export/import/update/delete가 요구사항대로 동작하는가?",
   ["8개 명령은 모두 argparse 하위 명령으로 등록되고, main()이 각 cmd_* 함수로 연결합니다.",
    "add는 입력을 검증해 파일 끝에 한 줄 추가, list는 최신순, search는 조건에 맞는 거래만 출력합니다."],
   [code("cli.py", 232, 235, 239, 248, 269, 279, 283, 287),
    code("cli.py", (297, 300))],
   [term("$ printf '2024-01-15\\nexpense\\nfood\\n15000\\n점심\\nmeal\\n' | python3 -m budget_app add --data-dir demo\n"
         "... [저장 완료] id=TX-000001\n"
         "$ printf '2024-01-25\\nincome\\nsalary\\n3000000\\n월급\\n\\n' | python3 -m budget_app add --data-dir demo\n"
         "... [저장 완료] id=TX-000002"),
    term("$ python3 -m budget_app list --data-dir demo\n"
         "TX-000002 | 2024-01-25 | income  | salary | 3000000 | 월급\n"
         "TX-000001 | 2024-01-15 | expense | food | 15000 | 점심 | #meal"),
    term("$ python3 -m budget_app search --category food --data-dir demo\n"
         "TX-000001 | 2024-01-15 | expense | food | 15000 | 점심 | #meal")])

qa("항목 1-1 | 8개 명령 동작 (2/2): summary · update · import · export · delete",
   "add/list/search/summary/export/import/update/delete가 요구사항대로 동작하는가?",
   ["summary는 월별 수입·지출·잔액·예산 사용률, update/delete는 id로 거래를 찾아 파일에 반영,",
    "import/export는 CSV로 거래를 주고받습니다. 모두 종료 코드 0으로 끝납니다."],
   [code("service.py", (230, 238)),
    code("service.py", (202, 207)),
    code("service.py", (210, 211))],
   [term("$ python3 -m budget_app budget set --month 2024-01 --amount 100000 --data-dir demo\n"
         "[저장 완료] 2024-01 예산 100000원\n"
         "$ python3 -m budget_app summary --month 2024-01 --data-dir demo\n"
         "총 수입: 3000000원\n총 지출: 15000원\n잔액: 2985000원\n"
         "예산: 100000원 (사용률 15.0%)\n\n지출 TOP 5\n1) food 15000원"),
    term("$ python3 -m budget_app update --id TX-000001 --amount 18000 --data-dir demo\n"
         "[수정 완료] id=TX-000001\n"
         "$ python3 -m budget_app import --from imp.csv --data-dir demo\n"
         "[완료] imported=2, skipped=1\n"
         "$ python3 -m budget_app export --out out.csv --month 2024-01 --data-dir demo\n"
         "[완료] out.csv (4 records)\n"
         "$ python3 -m budget_app delete --id TX-000002 --data-dir demo\n"
         "[삭제 완료] id=TX-000002")])

qa("항목 1-2 | 재실행 후 데이터 유지 (저장 파일 3개)",
   "프로그램 재실행 후에도 거래/카테고리/예산 데이터가 유지되는가? (저장 파일 3개 이상)",
   ["거래·카테고리·예산을 JSONL 파일 3개에 나누어 저장합니다. 명령마다 새 프로세스로 실행되므로,",
    "다음 명령에서 앞 명령의 결과가 보이는 것 자체가 재실행 후 유지된다는 근거입니다."],
   [code("cli.py", (16, 21)),
    code("repository.py", (40, 43))],
   [term("$ printf 'hobby\\n' | python3 -m budget_app category add --data-dir demo\n"
         "카테고리명: [저장 완료] category=hobby\n"
         "$ ls demo\n"
         "budgets.jsonl  categories.jsonl  transactions.jsonl"),
    term("$ python3 -m budget_app list --data-dir demo\n"
         "TX-000004 | 2024-01-22 | income  | salary | 50000 | 보너스\n"
         "TX-000003 | 2024-01-20 | expense | transport | 1250 | 지하철\n"
         "TX-000001 | 2024-01-15 | expense | food | 18000 | 점심 | #meal\n"
         "$ python3 -m budget_app category list --data-dir demo\n"
         "- food\n- transport\n- rent\n- salary\n- etc\n- hobby\n"
         "$ cat demo/budgets.jsonl\n"
         "{\"month\": \"2024-01\", \"amount\": 100000}")])

qa("항목 1-3 | category add/list/remove와 사용 중 카테고리 보호",
   "category add/list/remove가 정상 동작하는가? (삭제 시 사용 중인 카테고리 처리 포함)",
   ["remove는 먼저 거래 파일에서 사용 중인 카테고리 집합을 구해, 사용 중이면 AppError로 거부합니다.",
    "사용하지 않는 카테고리만 categories.jsonl에서 제거합니다. (add/list 실행은 앞 슬라이드 참고)"],
   [code("service.py", (88, 96)),
    code("repository.py", (118, 119)),
    code("repository.py", (150, 151))],
   [term("$ printf 'food\\n' | python3 -m budget_app category remove --data-dir demo\n"
         "삭제할 카테고리명: [오류] 'food' 카테고리를 사용 중인 거래가\n"
         "  있어 삭제할 수 없습니다.\n"
         "[힌트] 해당 카테고리를 사용하는 거래를 먼저 다른 카테고리로\n"
         "  옮기거나 삭제하세요.\n"
         "$ echo $?\n1", "사용 중인 food → 거부, 종료 코드 1"),
    term("$ printf 'hobby\\n' | python3 -m budget_app category remove --data-dir demo\n"
         "삭제할 카테고리명: [삭제 완료] category=hobby", "사용하지 않는 hobby → 삭제")])

qa("항목 1-4 | budget set 저장과 예산 사용률·초과 경고",
   "budget set이 저장되며, summary에서 예산 사용률/초과 여부가 출력되는가?",
   ["budget set은 budgets.jsonl에 월별로 저장(같은 월은 덮어씀)하고, summary가 지출÷예산으로",
    "사용률을 계산해 출력합니다. 지출이 예산보다 크면 [경고]를 추가로 출력합니다."],
   [code("repository.py", (167, 171)),
    code("service.py", (241, 246)),
    code("cli.py", (119, 122))],
   [term("$ python3 -m budget_app budget set --month 2024-01 --amount 10000 --data-dir demo\n"
         "[저장 완료] 2024-01 예산 10000원\n"
         "$ cat demo/budgets.jsonl\n"
         "{\"month\": \"2024-01\", \"amount\": 10000}"),
    term("$ python3 -m budget_app summary --month 2024-01 --data-dir demo\n"
         "총 수입: 50000원\n총 지출: 19250원\n잔액: 30750원\n"
         "예산: 10000원 (사용률 192.5%)\n[경고] 이번 달 예산을 초과했습니다.\n\n"
         "지출 TOP 5\n1) food 18000원\n2) transport 1250원")])

qa("항목 1-5 | import/export CSV 스키마 (UTF-8 · 헤더 · 컬럼)",
   "import/export가 명시된 CSV 스키마(UTF-8, 헤더, 컬럼)로 동작하는가?",
   ["컬럼은 CSV_FIELDS 한 곳에 정의(date,type,category,amount,memo,tags)하고, import는 UTF-8로 읽어",
    "헤더 이름으로 값을 꺼내며(DictReader), export는 같은 컬럼 순서로 헤더를 먼저 씁니다(DictWriter)."],
   [code("service.py", 18),
    code("service.py", (268, 270)),
    code("service.py", (310, 312)),
    code("service.py", 321)],
   [term("$ cat imp.csv\n"
         "date,type,category,amount,memo,tags\n"
         "2024-01-20,expense,transport,1250,지하철,\n"
         "2024-01-21,expense,없는카테고리,1000,깨짐,\n"
         "2024-01-22,income,salary,50000,보너스,", "가져온 파일"),
    term("$ cat out.csv\n"
         "date,type,category,amount,memo,tags\n"
         "2024-01-15,expense,food,18000,점심,meal\n"
         "2024-01-25,income,salary,3000000,월급,\n"
         "2024-01-20,expense,transport,1250,지하철,\n"
         "2024-01-22,income,salary,50000,보너스,", "내보낸 파일 (항목 1-1의 export 결과, 헤더 포함)")])

qa("항목 1-6·7 | 오류 메시지 · 해결 힌트 · 종료 코드",
   ["잘못된 입력/파일 오류에서 스택트레이스 없이 오류 메시지와 해결 힌트를 출력하는가?",
    "오류 상황에서 종료 코드가 0이 아님을 확인할 수 있는가?"],
   ["예상한 오류는 AppError(메시지+힌트)로 올리고, 모든 명령을 감싼 handle_errors가 이를 잡아",
    "표준에러에 [오류]/[힌트]만 출력하고 1을 반환합니다. __main__이 이 값을 sys.exit로 넘깁니다."],
   [code("decorators.py", (26, 34)),
    code("__main__.py", (6, 7))],
   [term("$ python3 -m budget_app update --id TX-999999 --amount 1000 --data-dir demo\n"
         "[오류] 존재하지 않는 거래 id입니다: TX-999999\n"
         "[힌트] `list` 명령으로 올바른 id를 확인하세요.\n"
         "$ echo $?\n1"),
    term("$ python3 -m budget_app export --out x.csv --data-dir demo\n"
         "[오류] export는 --month 또는 --from/--to 조건이 최소 1개 필요합니다.\n"
         "[힌트] 예: export --out export.csv --month 2024-01\n"
         "$ echo $?\n1"),
    term("$ python3 -m budget_app import --from nofile.csv --data-dir demo\n"
         "[오류] 파일을 찾을 수 없습니다: nofile.csv\n"
         "$ echo $?\n1\n"
         "$ python3 -m budget_app summary --month 2024-13 --data-dir demo\n"
         "[오류] 월 형식이 올바르지 않습니다 (YYYY-MM).\n"
         "[힌트] 예: 2024-01\n"
         "$ echo $?\n1")],
   note="참고: 파일을 찾을 수 없는 오류는 힌트 없이 메시지만 출력합니다 (service.py:264).")

# ---------------------------------------------------------------------------
# 항목 2
qa("항목 2-1 | 모듈 분리와 책임",
   "코드가 3개 이상 모듈로 분리되어 있고, 각 모듈의 책임을 “어떻게” 나눴는지 설명할 수 있는가?",
   ["바뀌는 이유가 다른 것끼리 나눴습니다: 화면(cli) · 규칙(service) · 파일(repository) · 데이터(models)",
    "· 공통 처리(decorators). import 방향이 cli → service → repository → models 한쪽으로만 흐릅니다."],
   [code("cli.py", (1, 3)),
    code("service.py", (1, 4)),
    code("repository.py", (1, 6)),
    code("models.py", 1)],
   [term("$ ls budget_app/*.py\n"
         "__init__.py  __main__.py  cli.py  decorators.py\n"
         "errors.py  models.py  repository.py  service.py"),
    term("$ grep -n \"^from budget_app\" budget_app/*.py\n"
         "__main__.py:4:   from budget_app.cli import main\n"
         "repository.py:16:from budget_app.models import Transaction\n"
         "decorators.py:11:from budget_app.errors import AppError\n"
         "service.py:13:   from budget_app.errors import AppError\n"
         "service.py:14:   from budget_app.models import ...\n"
         "service.py:15:   from budget_app.repository import ...\n"
         "cli.py:10:       from budget_app.decorators import ...\n"
         "cli.py:11:       from budget_app.errors import AppError\n"
         "cli.py:12:       from budget_app.repository import ...\n"
         "cli.py:13:       from budget_app.service import BudgetService",
         "의존 방향 확인 (repository는 service/cli를 import하지 않음)")])

qa("항목 2-2 | 클래스 책임 경계",
   "최소 2개 이상의 클래스에 부여한 책임 경계를 “어떻게” 정했는지 설명할 수 있는가?",
   ["Transaction=거래 한 건의 데이터, Repository/Store=파일 하나씩 입출력, BudgetService=검증과 규칙.",
    "Service는 저장소를 생성자로 받아 쓰기만 하고(파일 형식을 모름), 저장소는 업무 규칙을 모릅니다."],
   [code("models.py", (9, 19)),
    code("repository.py", (46, 47), (122, 123), (155, 156)),
    code("service.py", (68, 74))],
   [term("$ grep -n \"^class\" budget_app/*.py\n"
         "errors.py:5:       class AppError(Exception):\n"
         "service.py:68:     class BudgetService:\n"
         "models.py:10:      class Transaction:\n"
         "repository.py:46:  class TransactionRepository:\n"
         "repository.py:122: class CategoryStore:\n"
         "repository.py:155: class BudgetStore:"),
    term("경계 예시: 카테고리 삭제\n"
         "  규칙 판단 → BudgetService.remove_category (service.py:88)\n"
         "  사용 여부 → TransactionRepository.used_categories\n"
         "              (repository.py:118)\n"
         "  파일 수정 → CategoryStore.remove (repository.py:146)",
         "한 기능이 세 클래스에 어떻게 나뉘는가")])

qa("항목 2-3 | 파일 기반 update/delete의 안전한 처리",
   "파일 기반 update/delete를 “어떻게” 안전하게 처리했는지 설명할 수 있는가?",
   ["JSONL은 중간 행만 덮어쓸 수 없어 전체를 다시 씁니다. 원본을 직접 쓰지 않고 .tmp 파일에 먼저 쓴 뒤",
    "os.replace로 한 번에 교체하므로, 쓰는 도중 중단돼도 원본이 반쯤 쓰인 상태로 남지 않습니다."],
   [code("repository.py", (30, 37)),
    code("repository.py", 94, (97, 104))],
   [term("$ python3 -m budget_app delete --id TX-000003 --data-dir demo\n"
         "[삭제 완료] id=TX-000003\n"
         "$ ls demo\n"
         "budgets.jsonl  categories.jsonl  transactions.jsonl",
         "교체 후 .tmp 파일이 남지 않음"),
    term("$ python3 -m budget_app delete --id TX-999999 --data-dir demo\n"
         "[오류] 존재하지 않는 거래 id입니다: TX-999999\n"
         "[힌트] `list` 명령으로 올바른 id를 확인하세요.",
         "없는 id → found=False라 파일을 다시 쓰지 않음 (repository.py:114)"),
    term("한계: 전원 차단 등 저장장치 장애까지 보장하지는 않음\n"
         "      (fsync 미사용, 동시 실행 잠금 없음)")])

# ---------------------------------------------------------------------------
# 항목 3
qa("항목 3-1 | list/search 제너레이터 스트리밍",
   "list/search를 제너레이터로 스트리밍 처리한 방식을 “어떻게” 구현했고, “왜” 유리한지 설명할 수 있는가?",
   ["파일을 한 줄씩 읽어 yield하고, list는 deque(maxlen=N)로 최근 N건만 유지합니다. 거래가 늘어도",
    "메모리가 N건 분량으로 고정됩니다. 단, search는 최신순 정렬 때문에 일치한 결과는 리스트로 모읍니다."],
   [code("repository.py", (19, 27)),
    code("repository.py", (52, 55)),
    code("repository.py", (79, 82)),
    code("repository.py", (91, 92))],
   [term("$ python3 -c \"from pathlib import Path; from budget_app.repository import TransactionRepository as R; g = R(Path('demo/transactions.jsonl')).iter_all(); print(type(g).__name__); print(next(g))\"\n"
         "generator\n"
         "Transaction(id='TX-000001', date='2024-01-15', type='expense',\n"
         "  category='food', amount=18000, memo='점심', tags=['meal'])",
         "iter_all()은 제너레이터 — next()로 한 건씩 꺼냄"),
    term("$ python3 -m budget_app list --limit 1 --data-dir demo\n"
         "TX-000004 | 2024-01-22 | income  | salary | 50000 | 보너스"),
    term("거래 10만 건 파일(12.9MB)에서 최대 메모리 (/usr/bin/time -l)\n"
         "  list --limit 5          13 MB  (스트리밍)\n"
         "  search --category food  23 MB  (일치한 약 2만 건을 모음)\n"
         "  update                  79 MB  (전체 행을 리스트로 모음)",
         "왜 유리한가: 실측 비교 (항목 4-2 측정과 같은 데이터)")])

qa("항목 3-2 | 데코레이터로 분리한 공통 기능",
   "데코레이터로 분리한 공통 기능이 무엇이며, “왜” 분리가 필요했는지 설명할 수 있는가?",
   ["handle_errors(오류→메시지·힌트·종료 코드)와 log_execution(실행 로그·시간 측정)을 모든 cmd_* 함수에",
    "붙였습니다. 12개 명령에 같은 try/except와 로그 코드를 반복하지 않고, 명령 함수는 기능만 담습니다."],
   [code("decorators.py", (39, 58)),
    code("cli.py", (41, 43))],
   [term("$ grep -c \"^@handle_errors\" budget_app/cli.py\n12", "적용된 명령 함수 수"),
    term("$ tail -3 budget_app.log\n"
         "[2026-09-27 15:23:36] cmd_delete status=OK elapsed=0.7ms\n"
         "[2026-09-27 15:23:36] cmd_delete status=ERROR elapsed=0.3ms\n"
         "[2026-09-27 15:23:36] cmd_list status=OK elapsed=0.2ms",
         "log_execution 기록 (성공·실패 모두 남음, 시각은 실행 때마다 다름)"),
    term("$ python3 -c \"from budget_app.cli import cmd_add; print(cmd_add.__name__)\"\n"
         "cmd_add",
         "functools.wraps로 감싼 뒤에도 원래 이름 유지 → 로그에 명령명 기록")])

qa("항목 3-3 | 타입 힌트의 이점",
   "타입 힌트를 적용해 얻는 이점을 실제 코드 예로 “어떻게” 확인했고 “왜” 도움이 되는지 설명할 수 있는가?",
   ["함수 정의만 보고 입력·반환 타입을 알 수 있고(str→int, 생략 가능한 str | None, 제너레이터 반환),",
    "IDE 자동완성과 정적 분석에 쓰입니다. 단, 실행 중 값 검사는 하지 않으므로 검증은 validate_*가 합니다."],
   [code("service.py", (52, 59)),
    code("service.py", (134, 142)),
    code("repository.py", 52)],
   [term("$ python3 -c \"import typing; from budget_app.service import validate_amount, BudgetService; print(typing.get_type_hints(validate_amount)); print(typing.get_type_hints(BudgetService.search_transactions)['date_from'])\"\n"
         "{'amount_str': <class 'str'>, 'return': <class 'int'>}\n"
         "str | None",
         "코드에 적은 타입 정보를 실행 환경에서 그대로 읽을 수 있음"),
    term("$ python3 -m budget_app add --data-dir demo\n"
         "날짜(YYYY-MM-DD): 2024-01-15\n"
         "타입(income/expense): expense\n"
         "카테고리: food\n"
         "금액(양수): abc\n"
         "[오류] 금액은 숫자여야 합니다.\n"
         "[힌트] 예: 15000\n"
         "금액(양수):                       ← 다시 입력 요청",
         "잘못된 값은 타입 힌트가 아니라 validate_amount(service.py:53-56)가 막음")],
   note="mypy 같은 정적 분석 도구는 이 환경에 설치되어 있지 않아 실행하지 않았습니다.")

# ---------------------------------------------------------------------------
# 항목 4
qa("항목 4-1 | 저장 포맷: JSONL vs CSV",
   "JSONL과 CSV 중 선택한 저장 포맷의 장단점을 비교하고, “왜” 그 포맷을 택했는지 근거를 말할 수 있는가?",
   ["JSONL: 한 줄 = 한 거래라 추가는 끝에 한 줄 쓰기, tags 같은 목록도 그대로 저장 / 단점: 검색·수정은 전체 처리.",
    "CSV: 스프레드시트 호환이 장점이나 목록은 \"a,b\" 문자열로 바꿔야 함 → 내부 저장은 JSONL, 교환은 CSV."],
   [code("repository.py", (40, 43)),
    code("models.py", 19),
    code("service.py", 321),
    code("service.py", (62, 65))],
   [term("$ head -1 demo/transactions.jsonl\n"
         "{\"id\": \"TX-000001\", \"date\": \"2024-01-15\", \"type\": \"expense\",\n"
         " \"category\": \"food\", \"amount\": 18000, \"memo\": \"점심\",\n"
         " \"tags\": [\"meal\"]}",
         "JSONL: 숫자·목록 타입이 그대로 보존됨"),
    term("$ sed -n 2p out.csv\n"
         "2024-01-15,expense,food,18000,점심,meal",
         "CSV: 모든 값이 문자열, tags는 쉼표로 이어 붙인 문자열"),
    term("판단 근거\n"
         "  추가(add)가 가장 잦은 작업 → JSONL은 append 한 번\n"
         "  표준 라이브러리 json만으로 읽기/쓰기\n"
         "  사용자 교환은 CSV로 따로 제공 (import/export)")])

qa("항목 4-2 | 거래 10만 건에서의 병목과 개선",
   "거래가 10만 건으로 늘어난다면, 현재 구조에서 병목이 어디이며 “어떻게” 개선할지 설명할 수 있는가?",
   ["실측 결과 가장 큰 병목은 import입니다. 행마다 next_id()가 파일 전체를 다시 읽어 1,000행에 515초가",
    "걸렸습니다. 개선: ID 카운터 저장, import 일괄 쓰기, 날짜·카테고리 인덱스, 규모가 더 크면 SQLite."],
   [code("repository.py", (57, 65)),
    code("service.py", (118, 120), 128)],
   [term("command                    time      max memory\n"
         "list --limit 5             0.52 s     13 MB\n"
         "search --category food     0.57 s     23 MB\n"
         "summary --month 2024-03    0.54 s     13 MB\n"
         "export --month 2024-03     0.56 s     17 MB\n"
         "add (1 row)                0.59 s     13 MB\n"
         "update (1 row)             1.46 s     79 MB\n"
         "import (1,000 rows)      515.46 s     13 MB",
         "$ /usr/bin/time -l python3 -m budget_app <명령> --data-dir bench  (거래 10만 건, 12.9MB)"),
    term("병목 원인과 개선\n"
         "① import: 행마다 next_id()로 전체 읽기\n"
         "   → 비용 = 행 수 × 거래 수\n"
         "   → 마지막 id를 따로 저장하고 한 번에 append\n"
         "② 조회: 인덱스 없이 전체 순회 → 월별 파일 분할이나 인덱스\n"
         "③ update/delete: 전체를 메모리에 모아 재작성\n"
         "   → 한 줄씩 .tmp로 흘려 쓰기, 또는 SQLite로 이전")])

qa("항목 4-3 | 깨진 CSV 행과 사용자 신뢰",
   "import CSV에 일부 깨진 행이 섞이면, “어떻게” 처리해 사용자 신뢰를 지킬지(부분 성공/롤백/리포트) 설명할 수 있는가?",
   ["현재는 행 단위 부분 성공입니다. 행마다 검증해 실패(AppError·열 누락)한 행만 건너뛰고 개수를 알려줍니다.",
    "한계: 몇 번째 행이 왜 실패했는지 알려주지 않음 → 행 번호·사유 리포트, 필요하면 전체 검증 후 일괄 저장(롤백)."],
   [code("service.py", (266, 283))],
   [term("$ cat imp.csv\n"
         "date,type,category,amount,memo,tags\n"
         "2024-01-20,expense,transport,1250,지하철,\n"
         "2024-01-21,expense,없는카테고리,1000,깨짐,  ← 미등록 카테고리\n"
         "2024-01-22,income,salary,50000,보너스,\n"
         "$ python3 -m budget_app import --from imp.csv --data-dir demo\n"
         "[완료] imported=2, skipped=1\n"
         "$ echo $?\n0",
         "정상 행 2개는 저장, 깨진 행 1개는 건너뜀 (행 번호·사유는 출력 안 함)"),
    term("개선안\n"
         "  리포트: [건너뜀] 3행 category=없는카테고리 (미등록 카테고리)\n"
         "  롤백:   모든 행을 먼저 검증 → 하나라도 실패하면 저장 안 함,\n"
         "          전부 통과하면 한 번에 append\n"
         "  선택:   --strict 옵션으로 롤백/부분 성공을 사용자가 고름")])

# ---------------------------------------------------------------------------
# 마무리
slide, y = new_slide("마무리 | 발표용 요약 답변")
rect(slide, 0.72, y + 0.2, 11.9, 4.8, WHITE, BORDER)
text(slide, 0.95, y + 0.4, 11.4, 4.5,
     "거래·카테고리·예산을 JSONL 파일 세 개에 나누어 저장하고 (cli.py:16-21),\n"
     "CLI·서비스·저장소·모델로 책임을 분리했습니다.\n\n"
     "파일은 제너레이터로 한 줄씩 읽고 (repository.py:19-27), 목록은 최근 N건만 유지합니다 (repository.py:79-82).\n"
     "수정과 삭제는 임시 파일에 먼저 기록한 뒤 교체합니다 (repository.py:30-37).\n\n"
     "공통 오류 처리와 실행 로그는 데코레이터로 적용했고 (decorators.py), 타입 힌트로 입출력 계약을 드러냈습니다.\n\n"
     "한계도 확인했습니다. 검색 결과는 메모리에 모이고, CSV 가져오기는 행 단위 부분 성공이며\n"
     "실패 사유를 알려주지 않습니다. 10만 건 실측에서 import 1,000행이 515초 걸렸으므로\n"
     "ID 생성과 일괄 쓰기부터 개선하겠습니다.",
     15, TEXT)

with tempfile.TemporaryDirectory() as tmp:
    src = Path(tmp) / "평가자_답변.pptx"
    prs.save(src)
    convert(src, OUT)
print("saved", OUT, len(prs.slides), "slides")

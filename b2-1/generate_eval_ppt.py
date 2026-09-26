from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

OUT = "/Users/cspag5955/AI-SW-Basic/b2-1/평가질문_설명자료.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(18, 52, 98)
BLUE = RGBColor(24, 105, 182)
LIGHT = RGBColor(239, 246, 255)
WHITE = RGBColor(255, 255, 255)
TEXT = RGBColor(34, 34, 34)
MUTED = RGBColor(92, 102, 115)
GREEN = RGBColor(22, 140, 90)
PALE = RGBColor(230, 242, 255)


def add_footer(slide):
    box = slide.shapes.add_textbox(Inches(11.95), Inches(7.08), Inches(0.7), Inches(0.2))
    p = box.text_frame.paragraphs[0]
    p.text = f"{len(prs.slides):02d}"
    p.alignment = PP_ALIGN.RIGHT
    p.font.size = Pt(9)
    p.font.color.rgb = MUTED


def title(slide, text, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.72), Inches(0.38), Inches(11.9), Inches(0.55))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = NAVY
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.76), Inches(1.02), Inches(11.8), Inches(0.4))
        sp = sub.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(13)
        sp.font.color.rgb = MUTED


def new_slide(name, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT
    title(slide, name, subtitle)
    add_footer(slide)
    return slide


def card(slide, x, y, w, h, heading, body, color=BLUE, fill=WHITE, body_size=15):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = RGBColor(215, 226, 239)
    box = slide.shapes.add_textbox(Inches(x + 0.22), Inches(y + 0.17), Inches(w - 0.44), Inches(h - 0.3))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = heading
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = color
    p.space_after = Pt(9)
    for line in body.split("\n"):
        q = tf.add_paragraph()
        q.text = line
        q.font.size = Pt(body_size)
        q.font.color.rgb = TEXT
        q.space_after = Pt(4)
    return shape


def bullets(name, entries, subtitle=None):
    slide = new_slide(name, subtitle)
    y = 1.6
    for heading, body in entries:
        card(slide, 0.82, y, 11.68, 0.88, heading, body, body_size=14)
        y += 1.03
    return slide


def qa(name, question, answer, example=None):
    slide = new_slide(name)
    card(slide, 0.82, 1.55, 11.68, 1.22, "평가 질문", question, color=BLUE, fill=WHITE, body_size=16)
    card(slide, 0.82, 3.02, 11.68, 2.05, "답변 요점", answer, color=GREEN, fill=WHITE, body_size=16)
    if example:
        card(slide, 0.82, 5.35, 11.68, 0.92, "코드 / 근거", example, color=NAVY, fill=PALE, body_size=13)
    return slide


# 1. Cover
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
box = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.2), Inches(1.35))
p = box.text_frame.paragraphs[0]
p.text = "파일 기반 가계부 프로그램"
p.font.size = Pt(34); p.font.bold = True; p.font.color.rgb = WHITE
sub = slide.shapes.add_textbox(Inches(1.05), Inches(3.48), Inches(11), Inches(0.8))
sp = sub.text_frame.paragraphs[0]
sp.text = "평가 질문 설명 자료 | 구조 · 저장 · 검증 · 스트리밍"
sp.font.size = Pt(19); sp.font.color.rgb = RGBColor(205, 224, 247)
add_footer(slide)

# 2. Overview
slide = new_slide("1. 과제 결과물 한눈에 보기", "표준 라이브러리로 만든 Python CLI 가계부")
card(slide, 0.82, 1.65, 3.7, 3.75, "기능", "거래 추가·조회·검색\n월별 요약·예산\n카테고리 관리\n수정·삭제\nCSV 가져오기·내보내기")
card(slide, 4.82, 1.65, 3.7, 3.75, "저장", "JSONL 파일 3개\ntransactions.jsonl\ncategories.jsonl\nbudgets.jsonl\n앱을 종료해도 데이터 유지")
card(slide, 8.82, 1.65, 3.7, 3.75, "구조", "CLI → Service → Repository\nTransaction dataclass\n타입 힌트\n데코레이터 기반 공통 처리")

# 3. Data and CRUD
qa("2. 파일을 나누고 CRUD를 구현한 이유",
   "왜 거래·카테고리·예산을 세 파일로 나눴나요? 수정과 삭제는 어떻게 하나요?",
   "세 데이터는 성격과 규칙이 달라 별도 파일과 저장소 클래스로 관리합니다. 추가는 JSONL 끝에 한 줄을 덧붙입니다. 수정·삭제는 전체를 읽고 변경 결과를 임시 파일에 쓴 뒤 원본과 교체합니다.",
   "TransactionRepository.add / update / delete · CategoryStore · BudgetStore")

# 4. Architecture
slide = new_slide("3. 계층별 책임", "입출력, 규칙, 파일 형식을 서로 분리")
card(slide, 0.82, 1.6, 5.65, 1.55, "CLI · cli.py", "명령 옵션을 읽고 입력을 받아 결과를 출력")
card(slide, 6.82, 1.6, 5.65, 1.55, "Service · service.py", "입력 검증, 가계부 규칙, 월별 계산")
card(slide, 0.82, 3.5, 5.65, 1.55, "Repository · repository.py", "JSONL 파일 읽기·쓰기")
card(slide, 6.82, 3.5, 5.65, 1.55, "Model · models.py", "Transaction 데이터 필드와 변환")
arrow = slide.shapes.add_textbox(Inches(5.9), Inches(2.15), Inches(1.0), Inches(0.55))
ap = arrow.text_frame.paragraphs[0]; ap.text = "→"; ap.font.size = Pt(25); ap.font.bold = True; ap.font.color.rgb = BLUE
card(slide, 2.1, 5.55, 9.1, 0.75, "실행 흐름", "__main__.py가 cli.main()을 호출하고, CLI가 서비스와 저장소를 사용합니다.", body_size=14)

# 5. Streaming, accurately qualified
qa("4. 제너레이터와 스트리밍",
   "yield를 사용하면 어떤 점이 좋아지고, 모든 검색이 일정 메모리로 처리되나요?",
   "iter_all()은 JSONL을 한 줄씩 읽어 거래 객체를 yield합니다. list는 deque로 최근 N건만 유지합니다. 다만 search는 최신순 출력을 위해 조건에 맞는 결과를 모으며, CLI도 이를 list로 받습니다. 따라서 검색 결과가 많으면 메모리가 결과 수에 따라 늘어납니다.",
   "정확한 설명: 파일 읽기는 한 줄씩 진행하지만, 검색 결과 전체를 저장하지 않는 구조는 아닙니다.")

# 6. Atomic replace
qa("5. 원자적 쓰기",
   "원자적 쓰기란 무엇이며 update/delete에 왜 적용했나요?",
   "새 내용을 임시 파일에 모두 기록하고 os.replace()로 원본 파일과 교체하는 방식입니다. 기록 중 문제가 생기면 교체 단계에 도달하지 않아 원본을 그대로 둘 수 있어, 쓰다 만 내용으로 원본이 덮이는 위험을 줄입니다.",
   "현재 구현: 임시 파일은 원본과 같은 폴더에 생성한 뒤 os.replace() 호출")

# 7. Decorators
qa("6. 데코레이터로 공통 기능 분리",
   "@handle_errors와 @log_execution은 각각 어떤 역할을 하나요?",
   "log_execution은 명령 함수 실행 시간과 성공·실패를 로그에 기록합니다. handle_errors는 사용자 오류와 예외를 잡아 오류 메시지 및 종료 코드를 정합니다. 명령마다 같은 코드를 반복하지 않아도 됩니다.",
   "cli.py 명령 함수에 @handle_errors와 @log_execution 적용")

# 8. Validation and exit status
qa("7. 입력 검증과 오류 처리",
   "잘못된 날짜나 없는 거래 ID를 입력하면 어떻게 되나요?",
   "add의 날짜·타입·금액 입력은 유효할 때까지 다시 묻습니다. update/delete의 없는 ID 등은 AppError로 알리고, 데코레이터가 스택트레이스 대신 오류와 힌트를 출력한 뒤 0이 아닌 종료 코드를 반환합니다.",
   "근거 자료: 제출/results/13_add_invalid_date_retry.txt · 10_update_not_found.txt · 10_delete_not_found.txt")

# 9. CSV
qa("8. CSV 가져오기와 내보내기",
   "내부 저장은 JSONL인데 CSV 기능을 둔 이유는 무엇인가요?",
   "JSONL은 프로그램 내부 저장에 사용하고, CSV는 스프레드시트 등 다른 도구와 데이터를 주고받을 때 사용합니다. import는 행마다 검증하고 유효한 거래만 등록합니다. export는 거래 필드를 CSV 헤더에 맞춰 기록합니다.",
   "CSV 필드: date, type, category, amount, memo, tags")

# 10. Types and model
qa("9. dataclass와 타입 힌트",
   "Transaction을 dataclass로 만들고 타입 힌트를 쓴 이유는 무엇인가요?",
   "dataclass는 거래 한 건의 필드를 한 곳에 모으고 초기화·딕셔너리 변환을 간단하게 합니다. 타입 힌트는 함수가 어떤 값을 받고 돌려주는지 보여줘 코드를 읽고 변경하기 쉽게 합니다. 타입 힌트 자체가 실행 중 값을 자동 검증하지는 않습니다.",
   "예: validate_amount(amount_str: str) -> int")

# 11. Known constraints
bullets("10. 점검 결과와 설명할 때 주의할 점", [
    ("검색 메모리", "현재 search는 전체 일치 결과를 모아 역순으로 출력합니다. 대량 검색에서도 메모리 사용량이 일정하다고 말하지 않습니다."),
    ("내보내기 조건", "구현은 --month, --from, --to 중 하나만 있어도 허용합니다. 과제 문구가 시작일·종료일 쌍을 요구하는지 명확히 해 구현과 문서를 맞출 필요가 있습니다."),
    ("날짜 검증", "export의 --from/--to 값은 현재 날짜 형식 검증을 거치지 않습니다. 형식 오류도 검증하도록 보완하면 요구사항과 설명이 더 일치합니다."),
    ("원자성 설명", "원자적 교체는 중간 쓰기 실패에서 원본 보호를 돕습니다. 모든 종류의 저장장치 장애에서도 무조건 안전하다는 뜻으로 설명하지 않습니다."),
])

# 12. Summary
slide = new_slide("마무리: 짧게 설명하기")
card(slide, 0.95, 1.65, 11.45, 3.65, "발표 요약", "가계부 데이터는 거래·카테고리·예산 JSONL 파일로 저장합니다.\nCLI, 서비스, 저장소, 모델로 책임을 나눴습니다.\n파일은 제너레이터로 한 줄씩 읽고, 목록 조회는 최근 N건만 보관합니다.\n수정·삭제는 임시 파일 작성 후 교체해 원본 손상 위험을 줄입니다.\n공통 오류 처리와 실행 기록은 데코레이터로 적용했습니다.", body_size=17)

prs.save(OUT)
print(f"PPT created: {OUT}")

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

OUT = "/Users/cspag5955/AI-SW-Basic/b2-1/과제목표_기능요구사항_입증자료.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(18, 52, 98)
BLUE = RGBColor(24, 105, 182)
LIGHT = RGBColor(239, 246, 255)
WHITE = RGBColor(255, 255, 255)
TEXT = RGBColor(34, 34, 34)
MUTED = RGBColor(92, 102, 115)
GREEN = RGBColor(22, 140, 90)
PALE = RGBColor(230, 242, 255)
CODE_BG = RGBColor(30, 41, 59)
CODE_FG = RGBColor(226, 232, 240)


def footer(slide, note=None):
    if note:
        box = slide.shapes.add_textbox(Inches(0.72), Inches(7.04), Inches(10.6), Inches(0.25))
        p = box.text_frame.paragraphs[0]
        p.text = note
        p.font.size = Pt(8)
        p.font.color.rgb = MUTED
    box = slide.shapes.add_textbox(Inches(11.95), Inches(7.04), Inches(0.7), Inches(0.25))
    p = box.text_frame.paragraphs[0]
    p.text = f"{len(prs.slides):02d}"
    p.alignment = PP_ALIGN.RIGHT
    p.font.size = Pt(9)
    p.font.color.rgb = MUTED


def new_slide(heading, subtitle=None, note=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT
    box = slide.shapes.add_textbox(Inches(0.72), Inches(0.36), Inches(12.0), Inches(0.55))
    p = box.text_frame.paragraphs[0]
    p.text = heading
    p.font.size = Pt(25)
    p.font.bold = True
    p.font.color.rgb = NAVY
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.76), Inches(0.98), Inches(11.8), Inches(0.32))
        sp = sub.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(12)
        sp.font.color.rgb = MUTED
    footer(slide, note)
    return slide


def card(slide, x, y, w, h, heading, body, accent=BLUE, fill=WHITE, font_size=15):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = RGBColor(213, 225, 239)
    box = slide.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.13), Inches(w - 0.4), Inches(h - 0.23))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = heading
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = accent
    p.space_after = Pt(7)
    for line in body.split("\n"):
        q = tf.add_paragraph()
        q.text = line
        q.font.size = Pt(font_size)
        q.font.color.rgb = TEXT
        q.space_after = Pt(3)
    return shape


def codebox(slide, x, y, w, h, heading, body, font_size=12):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CODE_BG
    shape.line.color.rgb = CODE_BG
    box = slide.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.12), Inches(w - 0.36), Inches(h - 0.2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = heading
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor(125, 211, 252)
    p.space_after = Pt(5)
    for line in body.split("\n"):
        q = tf.add_paragraph()
        q.text = line
        q.font.name = "Consolas"
        q.font.size = Pt(font_size)
        q.font.color.rgb = CODE_FG
        q.space_after = Pt(2)
    return shape


def tile(slide, x, y, w, h, label, accent=BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = RGBColor(213, 225, 239)
    box = slide.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.08), Inches(w - 0.36), Inches(h - 0.16))
    p = box.text_frame.paragraphs[0]
    p.text = label
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = accent
    p.alignment = PP_ALIGN.CENTER


def goal_slide(num, heading, question, explanation, code_heading, code, evidence, note=None):
    slide = new_slide(f"과제 목표 {num} | {heading}", question, note)
    card(slide, 0.78, 1.48, 5.75, 2.0, "쉽게 설명하면", explanation, accent=GREEN, font_size=15)
    codebox(slide, 6.8, 1.48, 5.75, 2.0, code_heading, code, font_size=11)
    card(slide, 0.78, 3.82, 11.77, 1.68, "설명할 핵심", evidence, accent=BLUE, font_size=15)
    return slide


# Cover
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
b = slide.shapes.add_textbox(Inches(0.95), Inches(1.75), Inches(11.4), Inches(1.55))
p = b.text_frame.paragraphs[0]
p.text = "파일 기반 가계부 콘솔 프로그램"
p.font.size = Pt(33); p.font.bold = True; p.font.color.rgb = WHITE
s = slide.shapes.add_textbox(Inches(1.0), Inches(3.35), Inches(11.2), Inches(0.7))
sp = s.text_frame.paragraphs[0]
sp.text = "과제 목표 설명 · 기능 요구사항 충족 근거 · 실제 실행 결과"
sp.font.size = Pt(19); sp.font.color.rgb = RGBColor(205, 224, 247)
footer(slide, "B2-1 Python과 Git 심화")

# Roadmap
slide = new_slide("발표 구성", "목표를 설명하고, 요구사항별 구현과 실행 결과로 확인합니다.")
card(slide, 0.8, 1.55, 3.75, 3.95, "과제 목표 5가지", "1. 파일 저장과 CRUD\n2. 계층별 모듈 구조\n3. yield 스트리밍\n4. 데코레이터\n5. 타입 힌트", accent=BLUE, font_size=17)
card(slide, 4.8, 1.55, 3.75, 3.95, "요구사항 충족", "10개 CLI 기능\n3종 JSONL 파일\n입력 검증과 오류 처리\n수정·삭제 안전성\nCSV 가져오기/내보내기", accent=GREEN, font_size=17)
card(slide, 8.8, 1.55, 3.75, 3.95, "입증 자료", "실행 결과 텍스트\n저장 파일 예시\n코드 위치와 함수\n성공·실패 사례\n월 예산 초과 사례", accent=NAVY, font_size=17)

# Goal 1
goal_slide(1, "파일 저장과 CRUD", "파일 기반 저장으로 데이터를 계속 보관하고 기능을 구현할 수 있는가?",
           "가계부 데이터는 프로그램 메모리만이 아니라 JSONL 파일에 기록됩니다. 명령을 다시 실행해도 파일에서 읽기 때문에 내용이 남습니다.",
           "저장 파일 예시", '{"id":"TX-000001","date":"2024-01-15",\n "type":"expense","category":"food",\n "amount":18000,"memo":"점심(수정)"}',
           "추가(add)는 한 줄을 덧붙이고, 조회(list/search)는 파일을 읽습니다. 수정(update)과 삭제(delete)는 대상 거래를 반영한 새 파일을 작성해 교체합니다. 요약은 해당 월을 집계하고 CSV 입출력은 외부 교환 형식으로 변환합니다.",
           "근거: budget_app/repository.py · 제출/data_예시/transactions.jsonl")

# Goal 2
slide = new_slide("과제 목표 2 | 클래스와 모듈로 책임 나누기", "코드를 기능별 계층으로 나누면 각 파일의 역할을 설명하고 바꾸기 쉽습니다.")
card(slide, 0.8, 1.55, 5.65, 1.45, "진입점 · __main__.py", "cli.main()을 호출해 프로그램을 시작", font_size=14)
card(slide, 6.8, 1.55, 5.65, 1.45, "CLI · cli.py", "명령 옵션 파싱, 대화형 입력, 결과 출력", font_size=14)
card(slide, 0.8, 3.32, 5.65, 1.45, "서비스 · service.py", "날짜·금액 검증, 카테고리 규칙, 월별 계산", accent=GREEN, font_size=14)
card(slide, 6.8, 3.32, 5.65, 1.45, "저장소 · repository.py", "JSONL 파일 읽기·쓰기, 원자적 교체", accent=GREEN, font_size=14)
card(slide, 2.1, 5.12, 9.15, 0.92, "모델 · models.py", "Transaction dataclass로 거래 1건의 필드와 변환 방식을 정의", accent=NAVY, font_size=14)

# Goal 3
goal_slide(3, "yield 기반 제너레이터", "파일이 커져도 한 번에 전부 읽지 않도록 처리 흐름을 만들 수 있는가?",
           "yield는 값을 하나씩 만들어 호출한 쪽에 전달합니다. iter_all()은 JSONL의 한 줄을 읽어 Transaction 하나를 반환합니다.",
           "repository.py 핵심", 'def iter_all(self):\n    for row in _read_jsonl(self.path):\n        yield Transaction.from_dict(row)\n\ndef latest(self, limit):\n    window = deque(maxlen=limit)',
           "list는 deque에 최근 N건만 보관합니다. 반면 search는 최신순으로 반환하기 위해 현재 일치 결과를 모읍니다. 따라서 파일 읽기는 한 줄씩이어도 검색 결과가 많으면 메모리 사용량은 결과 수에 따라 증가합니다.",
           "근거: repository.py의 iter_all/latest/search · 제출/results/03_list.txt")

# Goal 4
goal_slide(4, "데코레이터", "로그·예외 처리·시간 측정을 명령 함수마다 반복하지 않고 분리할 수 있는가?",
           "데코레이터는 명령 함수를 감싸 공통 동작을 더합니다. 명령 본문은 기능에 집중하고, 로그와 오류 표시는 공통 코드로 처리합니다.",
           "cli.py 적용 예", '@handle_errors\n@log_execution\ndef cmd_add(args):\n    ...\n\n# log: status, elapsed ms',
           "log_execution은 함수명, 성공/실패, 실행 시간을 budget_app.log에 기록합니다. handle_errors는 AppError나 예외를 사용자용 오류/힌트로 바꾸고 정상은 0, 오류는 1을 반환합니다.",
           "근거: budget_app/decorators.py · 제출/budget_app_예시.log")

# Goal 5
goal_slide(5, "타입 힌트", "함수와 데이터가 주고받는 값의 종류를 명확히 설명할 수 있는가?",
           "타입 힌트는 코드의 입출력 설명서입니다. 에디터와 코드를 읽는 사람이 예상 값과 반환 값을 파악하는 데 도움을 줍니다. 런타임 자동 검증 기능과는 다릅니다.",
           "코드 예시", 'def validate_amount(amount_str: str) -> int:\n    ...\n\n@dataclass\nclass Transaction:\n    id: str\n    amount: int\n    tags: list[str]',
           "문자열 금액을 받아 양의 정수로 반환하는 계약과 거래 모델의 필드를 코드에서 바로 확인할 수 있습니다. 선택 값은 str | None처럼 표시합니다.",
           "근거: budget_app/service.py · budget_app/models.py")

# Requirements overview
slide = new_slide("기능 요구사항 | 10개 명령", "최종 결과물에 명시된 명령 단위로 정리했습니다.")
commands = ["add · 거래 추가", "list · 거래 목록", "search · 조건 검색", "summary · 월별 요약",
            "budget · 예산 설정", "category · 카테고리 관리", "update · 거래 수정",
            "delete · 거래 삭제", "import · CSV 가져오기", "export · CSV 내보내기"]
for i, label in enumerate(commands):
    col = i % 2
    row = i // 2
    x = 0.82 + col * 6.05
    y = 1.45 + row * 0.88
    tile(slide, x, y, 5.62, 0.69, f"{i + 1:02d}   {label}", accent=BLUE if col == 0 else GREEN)
card(slide, 1.35, 6.15, 10.6, 0.55, "공통 요구", "명령 도움말(--help) · 입력 검증 · 오류 원인/힌트 · 종료 코드", accent=NAVY, font_size=13)

# Persistence and data model
slide = new_slide("요구사항 | 영구 저장과 데이터 모델", "세 종류 데이터를 각각 JSONL로 저장하고 Transaction을 dataclass로 표현합니다.")
card(slide, 0.8, 1.55, 3.75, 2.6, "transactions.jsonl", "거래 ID\n날짜 / 수입·지출\n카테고리 / 금액\n메모 / 태그", font_size=15)
card(slide, 4.8, 1.55, 3.75, 2.6, "categories.jsonl", "등록된 카테고리 이름\n기본값: food, transport,\nrent, salary, etc", accent=GREEN, font_size=15)
card(slide, 8.8, 1.55, 3.75, 2.6, "budgets.jsonl", "월별 예산\nmonth + amount\n예: 2024-01 / 100000", accent=NAVY, font_size=15)
codebox(slide, 1.4, 4.55, 10.55, 1.35, "저장 파일 실제 예시", '{"month":"2024-01","amount":100000}\n한 줄마다 JSON 객체 하나를 저장하는 JSONL 형식', font_size=13)

# CRUD evidence
slide = new_slide("입증자료 | 거래 추가와 목록", "실제 결과 파일: 02_add_1.txt · 03_list.txt")
codebox(slide, 0.8, 1.55, 5.7, 3.75, "추가: python3 -m budget_app add", '날짜: 2024-01-15\n타입: expense\n카테고리: food\n금액: 15000\n메모: 점심\n태그: meal\n\n[저장 완료] id=TX-000001', font_size=14)
codebox(slide, 6.8, 1.55, 5.7, 3.75, "목록: python3 -m budget_app list --limit 3", 'TX-000004 | 2024-01-10 | expense | rent | 150000 | 월세\nTX-000003 | 2024-01-12 | expense | transport | 20000 | 택시\nTX-000002 | 2024-01-14 | income | salary | 3000000 | 월급\n\n종료 코드: 0', font_size=12)
card(slide, 1.5, 5.72, 10.3, 0.73, "확인", "추가는 ID를 발급해 저장하고, 목록은 최신 입력 순으로 최대 3건을 보여줍니다.", accent=GREEN, font_size=14)

# Search and validation evidence
slide = new_slide("입증자료 | 검색과 입력 검증", "실제 결과 파일: 04_search_category.txt · 13_add_invalid_date_retry.txt")
codebox(slide, 0.8, 1.55, 5.7, 3.7, "검색: --category food", 'TX-000001 | 2024-01-15 | expense | food | 15000 | 점심 | #meal\n\n종료 코드: 0', font_size=13)
codebox(slide, 6.8, 1.55, 5.7, 3.7, "잘못된 날짜 입력 후 재시도", '입력: 2024-13-40\n[오류] 날짜 형식이 올바르지 않습니다.\n[힌트] 예: 2024-01-15\n\n다시 입력: 2024-01-20\n[저장 완료] id=TX-000007\n종료 코드: 0', font_size=12)
card(slide, 1.5, 5.72, 10.3, 0.73, "확인", "검색 조건 적용과 대화형 입력 오류 복구 흐름을 실제 출력으로 확인했습니다.", accent=GREEN, font_size=14)

# Summary and budget evidence
slide = new_slide("입증자료 | 월별 요약과 예산 경고", "실제 결과 파일: 05_summary_before_budget.txt · 06_budget_set.txt · 07_summary_after_budget.txt")
codebox(slide, 0.8, 1.55, 5.7, 3.7, "예산 설정", '$ python3 -m budget_app budget set \\\n  --month 2024-01 --amount 100000\n[저장 완료] 2024-01 예산 100000원', font_size=14)
codebox(slide, 6.8, 1.55, 5.7, 3.7, "요약 결과", '총 수입: 3000000원\n총 지출: 185000원\n잔액: 2815000원\n예산: 100000원 (사용률 185.0%)\n[경고] 이번 달 예산을 초과했습니다.', font_size=14)
card(slide, 1.5, 5.72, 10.3, 0.73, "확인", "수입·지출·잔액·카테고리 TOP과 예산 사용률 및 초과 경고가 출력됩니다.", accent=GREEN, font_size=14)

# Category/update/delete proof
slide = new_slide("입증자료 | 카테고리 관리와 수정·삭제", "실제 결과 파일: 09_category_remove_blocked.txt · 10_update.txt · 10_delete_not_found.txt")
card(slide, 0.8, 1.55, 3.75, 3.35, "사용 중 카테고리 삭제", "food 카테고리 삭제 시도\n\n[오류] 사용 중인 거래가 있어 삭제할 수 없습니다.\n\n종료 코드: 1", accent=BLUE, font_size=13)
card(slide, 4.8, 1.55, 3.75, 3.35, "거래 수정", "update --id TX-000001\n--amount 18000\n--memo '점심(수정)'\n\n[수정 완료] id=TX-000001\n\n종료 코드: 0", accent=GREEN, font_size=13)
card(slide, 8.8, 1.55, 3.75, 3.35, "없는 거래 삭제", "delete --id TX-999999\n\n[오류] 존재하지 않는 거래 ID\n[힌트] list에서 ID 확인\n\n종료 코드: 1", accent=NAVY, font_size=13)
card(slide, 1.4, 5.35, 10.55, 0.82, "확인", "사용 중 카테고리 보호, 옵션 기반 수정, 존재하지 않는 ID 오류 처리를 확인했습니다.", accent=GREEN, font_size=14)

# Atomic write and errors
slide = new_slide("요구사항 | 원자적 교체와 오류 처리", "수정/삭제의 파일 안정성과 스택트레이스 없는 사용자 메시지")
card(slide, 0.8, 1.55, 5.7, 3.2, "원자적 쓰기 흐름", "① 임시 파일에 새 내용을 모두 기록\n\n② 기록이 끝나면 os.replace()로 교체\n\n쓰는 도중 실패하면 기존 파일을 유지할 수 있어, 절반만 기록된 내용으로 원본을 덮을 위험을 줄입니다.", accent=BLUE, font_size=14)
codebox(slide, 6.8, 1.55, 5.7, 3.2, "오류 처리 근거", 'AppError → [오류] 원인\n          → [힌트] 해결 방법\n          → 종료 코드 1\n\n정상 실행 → 종료 코드 0\n로그 → 함수명/status/소요시간', font_size=14)
card(slide, 1.5, 5.35, 10.3, 0.82, "근거", "budget_app/repository.py의 _atomic_write_jsonl · decorators.py의 handle_errors/log_execution", accent=GREEN, font_size=13)

# CSV proof
slide = new_slide("입증자료 | CSV 가져오기와 내보내기", "실제 결과 파일: 11_export.txt · 12_import.txt")
codebox(slide, 0.8, 1.55, 5.7, 3.45, "내보내기", '$ python3 -m budget_app export \\\n  --out export.csv --month 2024-01\n[완료] export.csv (3 records)\n\nCSV 헤더:\ndate,type,category,amount,memo,tags', font_size=13)
codebox(slide, 6.8, 1.55, 5.7, 3.45, "가져오기", '$ python3 -m budget_app import \\\n  --from import_test.csv\n\n[완료] imported=2, skipped=1\n\n정상 행은 등록하고, 잘못된 행은 건너뜁니다.', font_size=13)
card(slide, 1.5, 5.55, 10.3, 0.78, "CSV 최소 필드", "date · type · category · amount · memo(선택) · tags(선택) / UTF-8, 헤더 포함", accent=GREEN, font_size=14)

# Result examples
slide = new_slide("결과 예시 | 사용자 관점에서 보이는 동작", "명령을 실행하면 아래와 같이 결과를 확인할 수 있습니다.")
codebox(slide, 0.8, 1.55, 5.7, 3.65, "추가 후 성공 메시지", '$ python3 -m budget_app add\n날짜: 2024-01-15\n타입: expense\n카테고리: food\n금액: 15000\n\n[저장 완료] id=TX-000001', font_size=13)
codebox(slide, 6.8, 1.55, 5.7, 3.65, "조건 없는 export는 오류", '$ python3 -m budget_app export \\\n  --out export2.csv\n\n[오류] export는 --month 또는 --from/--to 조건이 최소 1개 필요합니다.\n[힌트] 예: export --out export.csv --month 2024-01\n\n종료 코드: 1', font_size=13)
card(slide, 1.5, 5.72, 10.3, 0.73, "결과", "정상 완료와 잘못된 사용 모두 메시지와 종료 코드로 구분됩니다.", accent=GREEN, font_size=14)

# Accuracy note
slide = new_slide("점검 메모 | 실행 결과와 요구사항을 정확히 설명하기", "실제 구현에서 보완이 필요한 부분도 확인합니다.")
card(slide, 0.82, 1.55, 11.65, 1.25, "검색의 메모리", "파일은 한 줄씩 읽지만, search는 최신순 출력을 위해 일치 결과를 모읍니다. 결과가 많으면 메모리 사용량도 늘어납니다.", accent=BLUE, font_size=14)
card(slide, 0.82, 3.05, 11.65, 1.25, "export 조건", "현재 구현은 --month, --from, --to 중 하나만 있어도 허용합니다. 과제 요구의 날짜 범위(--from와 --to)를 모두 요구하도록 보완할 필요가 있습니다.", accent=GREEN, font_size=14)
card(slide, 0.82, 4.55, 11.65, 1.25, "날짜 형식 검증", "export의 날짜 필터는 현재 YYYY-MM-DD 형식을 검증하지 않습니다. 엄격한 입력 검증을 위해 보완 항목으로 남깁니다.", accent=NAVY, font_size=14)

# Conclusion
slide = new_slide("마무리 | 요구사항과 근거를 연결해 설명하기")
card(slide, 0.9, 1.55, 11.5, 3.85, "발표 요약", "JSONL 파일 3개에 거래·카테고리·예산을 영구 저장합니다.\nCLI / Service / Repository / Model 계층으로 책임을 나눴습니다.\n제너레이터로 파일을 한 줄씩 읽고, 목록 조회는 최근 N건을 유지합니다.\n데코레이터가 오류 표시와 실행 로그·시간 측정을 공통 처리합니다.\n타입 힌트와 dataclass로 데이터 구조와 함수 계약을 읽기 쉽게 표현합니다.\n10개 명령의 성공·오류 실행 결과를 제출 자료로 확인할 수 있습니다.", accent=NAVY, font_size=17)

prs.save(OUT)
print(f"PPT created: {OUT}")

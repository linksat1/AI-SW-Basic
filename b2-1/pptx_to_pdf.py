"""python-pptx로 만든 PPT를 HTML로 재구성한 뒤 Chrome headless로 PDF 변환.

사용법: python3 pptx_to_pdf.py 평가/평가자_답변.pptx
PowerPoint/LibreOffice 없이 변환하기 위한 용도로, 이 폴더의 생성 스크립트가 쓰는
도형(텍스트 상자, 둥근 사각형, 표)만 지원한다.
"""
from __future__ import annotations

import html
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu

CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
SANS = '"Apple SD Gothic Neo", "Malgun Gothic", sans-serif'


def inch(v) -> str:
    return f"{Emu(v).inches:.3f}in"


def rgb(c) -> str:
    return f"#{c}"


def fill_color(fill):
    if fill.type == MSO_FILL.SOLID:
        return rgb(fill.fore_color.rgb)
    return None


def para_html(p, nowrap: bool) -> str:
    f = p.font
    style = [f"font-size:{f.size.pt if f.size else 18}pt", "line-height:1.2", "margin:0"]
    if f.bold:
        style.append("font-weight:700")
    try:
        if f.color and f.color.type is not None:
            style.append(f"color:{rgb(f.color.rgb)}")
    except AttributeError:
        pass
    if f.name == "Menlo":
        style.append('font-family:Menlo, monospace')
        style.append("white-space:pre-wrap")
    elif nowrap:
        style.append("white-space:pre")
    if p.space_after is not None:
        style.append(f"margin-bottom:{p.space_after.pt}pt")
    if p.alignment == PP_ALIGN.RIGHT:
        style.append("text-align:right")
    text = html.escape(p.text) or "&nbsp;"
    return f'<p style="{";".join(style)}">{text}</p>'


def shape_html(sh) -> str:
    pos = f"left:{inch(sh.left)};top:{inch(sh.top)};width:{inch(sh.width)};height:{inch(sh.height)}"
    if sh.has_table:
        tbl = sh.table
        cols = "".join(f'<col style="width:{inch(c.width)}">' for c in tbl.columns)
        rows = []
        for row in tbl.rows:
            cells = []
            for cell in row.cells:
                bg = fill_color(cell.fill) or "transparent"
                inner = "".join(para_html(p, False) for p in cell.text_frame.paragraphs)
                cells.append(f'<td style="background:{bg}">{inner}</td>')
            rows.append("<tr>" + "".join(cells) + "</tr>")
        return (f'<table class="tbl" style="{pos};height:auto"><colgroup>{cols}</colgroup>'
                + "".join(rows) + "</table>")

    parts = [f'<div class="shape" style="{pos}']
    if not sh.has_text_frame or not sh.text_frame.text.strip():
        # 도형(카드 배경)
        bg = fill_color(sh.fill) if hasattr(sh, "fill") else None
        if bg:
            parts.append(f";background:{bg}")
        try:
            if sh.line.fill.type == MSO_FILL.SOLID:
                parts.append(f";border:1px solid {rgb(sh.line.color.rgb)}")
        except (AttributeError, TypeError):
            pass
        if getattr(sh, "auto_shape_type", None) is not None:
            parts.append(";border-radius:0.12in")
        parts.append('"></div>')
        return "".join(parts)

    tf = sh.text_frame
    nowrap = tf.word_wrap is not True
    parts.append(';padding:0.05in 0.1in;box-sizing:border-box">')
    parts.extend(para_html(p, nowrap) for p in tf.paragraphs)
    parts.append("</div>")
    return "".join(parts)


def convert(src: Path, out: Path | None = None) -> Path:
    prs = Presentation(str(src))
    w, h = inch(prs.slide_width), inch(prs.slide_height)
    slides = []
    for slide in prs.slides:
        bg = fill_color(slide.background.fill) or "#ffffff"
        body = "".join(shape_html(sh) for sh in slide.shapes)
        slides.append(f'<section class="slide" style="background:{bg}">{body}</section>')
    doc = f"""<!doctype html><html><head><meta charset="utf-8"><style>
@page {{ size: {w} {h}; margin: 0; }}
* {{ -webkit-print-color-adjust: exact; print-color-adjust: exact; }}
body {{ margin: 0; font-family: {SANS}; }}
.slide {{ position: relative; width: {w}; height: {h}; overflow: hidden; page-break-after: always; }}
.shape, .tbl {{ position: absolute; }}
.tbl {{ border-collapse: collapse; table-layout: fixed; }}
.tbl td {{ padding: 0.05in 0.1in; vertical-align: top; border: 1px solid #d7e2ef; }}
</style></head><body>{"".join(slides)}</body></html>"""

    out = out or src.with_suffix(".pdf")
    with tempfile.TemporaryDirectory() as tmp:
        page = Path(tmp) / "slides.html"
        page.write_text(doc, encoding="utf-8")
        subprocess.run([CHROME, "--headless", "--disable-gpu", "--no-pdf-header-footer",
                        f"--print-to-pdf={out}", page.as_uri()],
                       check=True, capture_output=True)
    return out


if __name__ == "__main__":
    print("saved", convert(Path(sys.argv[1]).resolve()))
